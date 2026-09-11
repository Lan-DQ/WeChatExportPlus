# -*- coding: utf-8 -*-
"""Excel 导出器。

[WeChatExportPlus 改动] 改用 message_content.prepare() 的统一解析结果：
  - 图片/语音/链接/文件/系统消息不再被过滤掉；
  - 「类型」列输出可读中文（文字/图片/语音/链接/…）而不是 '类型49'。
列结构与样式保持不变（时间/发送者/消息内容/类型）。
"""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import message_content

HEADERS = ['时间', '发送者', '消息内容', '类型']
WIDTHS = [20, 18, 60, 10]


def export(msgs, path, my_name='我', title=None, session_title=None, prompt=None):
    wb = Workbook()
    ws = wb.active
    ws.title = '聊天记录'
    for i, w in enumerate(WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    hf = Font(bold=True, color='FFFFFF', size=11)
    hfill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
    ha = Alignment(horizontal='center', vertical='center')
    bd = Border(left=Side(style='thin', color='E5E7EB'),
                right=Side(style='thin', color='E5E7EB'),
                top=Side(style='thin', color='E5E7EB'),
                bottom=Side(style='thin', color='E5E7EB'))

    # 给 AI 的指令放在最前面几行：无论从哪开始读都先看到要求。
    # 表头与筛选范围随之整体下移，避免指令行被当成数据行。
    lead = 0
    if prompt and prompt.strip():
        ws.append(['# --- 给 AI 的指令（以下不是聊天记录）---'])
        lead = 1
        for line in prompt.strip().split('\n'):
            ws.append([line])
            lead += 1
        ws.append([])
        lead += 1
        for r in range(1, lead + 1):
            cell = ws.cell(row=r, column=1)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.font = Font(size=9, color='7F1D1D')
        ws.append([])
        lead += 1

    ws.append(HEADERS)
    header_row = ws.max_row
    for c in ws[header_row]:
        c.font = hf
        c.fill = hfill
        c.alignment = ha
        c.border = bd

    for m in msgs:
        text = (m.get('text') or '').strip()
        if not text:
            continue
        ws.append([
            m.get('time_str', ''),
            m.get('sender_display', ''),
            text,
            message_content.kind_label(m.get('kind', 'unknown')),
        ])
        r = ws.max_row
        for ci in range(1, 5):
            cell = ws.cell(row=r, column=ci)
            cell.alignment = Alignment(vertical='top', wrap_text=(ci == 3))
            cell.border = bd

    ws.freeze_panes = f'A{header_row + 1}'
    if ws.max_row > header_row:
        ws.auto_filter.ref = f'A{header_row}:D{ws.max_row}'
    wb.save(path)
    return path
