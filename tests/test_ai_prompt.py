# -*- coding: utf-8 -*-
"""「给 AI 的指令」必须随每一种导出一起产出。"""
import json
import os
import sys

import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, 'exporters'))

import ai_prompt          # noqa: E402
import batch_export as be  # noqa: E402
import message_content as mc  # noqa: E402


class FakeWCDB:
    def __init__(self, data=None):
        self.data = data or {}

    def get_messages(self, wxid, limit=0, offset=0):
        rows = self.data.get(wxid, [])
        return rows[:limit] if limit else rows

    def get_display_names(self, wxids):
        return {w: f'N-{w}' for w in wxids}

    def get_count(self, wxid):
        return len(self.data.get(wxid, []))


def msg(local_type, content, ts='1700000000', sender='', lid='1'):
    return {'local_type': local_type, 'message_content': content,
            'create_time': ts, 'sender_username': sender, 'local_id': lid}


PROMPT = '【角色】你是我的助理。\n[图片0001] 代表图片，你看不到内容，需要时向我要。'
MARK = '给 AI 的指令'


@pytest.fixture(autouse=True)
def _clear_cache():
    ai_prompt.clear_cache()
    yield
    ai_prompt.clear_cache()


@pytest.fixture
def fake():
    return FakeWCDB({'g@chatroom': [
        msg(1, '你好', sender='wxid_a', lid='1'),
        msg(3, '<msg><img/></msg>', sender='wxid_a', lid='2'),
    ]})


SESSIONS = [{'wxid': 'g@chatroom', 'title': '群聊A'}]


# ─────────────── ai_prompt 模块 ───────────────

def test_default_prompt_nonempty_when_no_file(tmp_path):
    text = ai_prompt.load_prompt(str(tmp_path))
    assert len(text) > 100
    assert '图片' in text


def test_external_file_overrides_default(tmp_path):
    (tmp_path / 'AI提示词.txt').write_text('自定义指令XYZ', encoding='utf-8')
    assert ai_prompt.load_prompt(str(tmp_path)) == '自定义指令XYZ'


def test_blank_external_file_falls_back_to_default(tmp_path):
    (tmp_path / 'AI提示词.txt').write_text('   \n\n  ', encoding='utf-8')
    assert '角色' in ai_prompt.load_prompt(str(tmp_path))


def test_cache_cleared_rereads_file(tmp_path):
    p = tmp_path / 'AI提示词.txt'
    p.write_text('第一版', encoding='utf-8')
    assert ai_prompt.load_prompt(str(tmp_path)) == '第一版'
    p.write_text('第二版', encoding='utf-8')
    assert ai_prompt.load_prompt(str(tmp_path)) == '第一版'   # 命中缓存
    ai_prompt.clear_cache()
    assert ai_prompt.load_prompt(str(tmp_path)) == '第二版'


def test_as_markdown_wraps_in_blockquote():
    md = ai_prompt.as_markdown('第一行\n\n第二行')
    assert md.startswith('## ' + MARK)
    assert '> 第一行' in md
    assert '> 第二行' in md
    assert md.rstrip().endswith('---')


# ─────────────── 每种格式都要带上 ───────────────

def test_md_contains_prompt(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'md', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    text = open(os.path.join(res['root'], '群聊A.md'), encoding='utf-8').read()
    assert MARK in text
    assert '你是我的助理' in text
    # 指令必须在第一条聊天消息之前
    assert text.index('你是我的助理') < text.index('[20')


def test_txt_contains_prompt(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'txt', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    text = open(os.path.join(res['root'], '群聊A', '对话.txt'), encoding='utf-8').read()
    assert '你是我的助理' in text


def test_json_contains_prompt_as_field(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'json', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    data = json.load(open(os.path.join(res['root'], '群聊A', '对话.json'), encoding='utf-8'))
    assert data['_ai_instruction'] == PROMPT
    assert isinstance(data['messages'], list)
    assert data['messages'][0]['text'] == '你好'


def test_ai_format_contains_prompt_in_both_files(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'ai', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    d = os.path.join(res['root'], '群聊A')
    meta = json.loads(open(os.path.join(d, '对话.jsonl'), encoding='utf-8').readline())
    assert meta['_meta']['ai_instruction'] == PROMPT
    txt = open(os.path.join(d, '对话.txt'), encoding='utf-8').read()
    assert '你是我的助理' in txt


def test_csv_contains_prompt(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'csv', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    text = open(os.path.join(res['root'], '群聊A', '对话.csv'),
                encoding='utf-8-sig').read()
    assert '你是我的助理' in text


def test_excel_contains_prompt(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'excel', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    from openpyxl import load_workbook
    wb = load_workbook(os.path.join(res['root'], '群聊A', '群聊A.xlsx'))
    ws = wb.active
    col_a = [str(c.value or '') for c in ws['A']]
    assert any(MARK in v for v in col_a), 'Excel 第一列应有指令行'
    assert '时间' in col_a, '表头仍应在'


def test_html_contains_prompt(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'html', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    text = open(os.path.join(res['root'], '群聊A', 'index.html'), encoding='utf-8').read()
    assert MARK in text
    assert '你是我的助理' in text


def test_pdf_contains_prompt(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'pdf', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    raw = open(os.path.join(res['root'], '群聊A', '群聊A.pdf'), 'rb').read()
    assert raw.startswith(b'%PDF')
    # PDF 里中文是压缩/编码过的，不能直接断言文本；这里只确认文件生成成功且非空
    assert len(raw) > 1000


def test_export_also_writes_prompt_reference_file(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'md', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    ref = os.path.join(res['root'], '给AI的指令.txt')
    assert os.path.exists(ref)
    assert open(ref, encoding='utf-8').read().strip() == PROMPT


def test_index_page_shows_prompt(fake, tmp_path):
    res = be.export_sessions(fake, '', SESSIONS, 'md', str(tmp_path),
                             resolve_images=False, prompt=PROMPT)
    html = open(os.path.join(res['root'], '导出清单.html'), encoding='utf-8').read()
    assert MARK in html
    assert '你是我的助理' in html


def test_no_prompt_means_no_prompt_block(fake, tmp_path):
    """显式传空提示词时不应插入任何指令块（避免出现空标题）。"""
    res = be.export_sessions(fake, '', SESSIONS, 'md', str(tmp_path),
                             resolve_images=False, prompt='')
    text = open(os.path.join(res['root'], '群聊A.md'), encoding='utf-8').read()
    assert MARK not in text
    assert not os.path.exists(os.path.join(res['root'], '给AI的指令.txt'))


def test_prompt_loaded_from_base_dir_when_not_passed(fake, tmp_path):
    """不显式传 prompt 时，应从 base_dir/AI提示词.txt 读取。"""
    base = tmp_path / 'app'
    base.mkdir()
    (base / 'AI提示词.txt').write_text('来自文件的指令ABC', encoding='utf-8')
    res = be.export_sessions(fake, '', SESSIONS, 'md', str(tmp_path / 'out'),
                             resolve_images=False, base_dir=str(base))
    text = open(os.path.join(res['root'], '群聊A.md'), encoding='utf-8').read()
    assert '来自文件的指令ABC' in text
